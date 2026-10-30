import hashlib
import re
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.errors import AppError
from app.models import (
    ConstructionDirection,
    ConstructionPlan,
    DictIndicatorType,
    DictIndicatorUnit,
    IndicatorAnnualTarget,
    OutputIndicator,
    SourceImportFile,
    SourceImportRow,
    SysOrgAlias,
    SysOrgUnit,
    SysUser,
)
from app.services.audit import write_log
from app.services.values import assign_target_fields, parse_decimal_text, parse_target, raw_text
from app.utils import model_dict, utcnow


SOURCE_FIELDS = [
    "序号",
    "建设任务/主要建设方向",
    "关键产出指标",
    "指标类型",
    "指标单位",
    "专业群5年量化目标",
    "专业群2026年量化目标",
    "专业群五年指标分配",
    "专业群2026年指标分配",
    "学校5年量化目标",
    "学校2026年量化目标",
    "学校2026年任务分工",
    "学校2027年任务分工",
    "学校2028年任务分工",
    "学校2029年任务分工",
    "学校2030年任务分工",
]
ROW_FIELD_NAMES = [
    "source_seq_raw",
    "direction_raw",
    "indicator_raw",
    "indicator_type_raw",
    "indicator_unit_raw",
    "professional_five_year_raw",
    "professional_2026_raw",
    "professional_five_year_allocation_raw",
    "professional_2026_allocation_raw",
    "school_five_year_raw",
    "school_2026_raw",
    "school_2026_division_raw",
    "school_2027_division_raw",
    "school_2028_division_raw",
    "school_2029_division_raw",
    "school_2030_division_raw",
]


SOURCE_FORMATS = {"MERGED", "SCHOOL", "TASK"}


def _find_header(sheet, source_format: str = "MERGED") -> tuple[int, list[str]] | None:
    for row_no in range(1, min(sheet.max_row, 30) + 1):
        values = [raw_text(sheet.cell(row_no, column).value) for column in range(1, 17)]
        if source_format == "MERGED" and values == SOURCE_FIELDS:
            return row_no, values
        if source_format == "SCHOOL" and values[1:3] == ["主要建设方向", "关键产出指标"]:
            second = [raw_text(sheet.cell(row_no + 1, column).value) for column in range(4, 9)]
            if second == ["指标类型", "指标单位", "5 年的量化目标", "2026年的量化目标", "2026年"]:
                return row_no + 1, values
        if source_format == "TASK" and values[:9] == ["序号", "建设任务", "关键产出指标", "指标类型", "指标单位", "5年量化目标", "2026年量化目标", "五年指标分配", "2026年指标分配"]:
            return row_no, values
    return None


def _cell_value(sheet, row_no: int, column: int):
    value = sheet.cell(row_no, column).value
    if value is not None:
        return value
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row_no <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return sheet.cell(merged.min_row, merged.min_col).value
    return None


def _source_values(sheet, row_no: int, source_format: str) -> list[Any]:
    if source_format == "MERGED":
        return [sheet.cell(row_no, column).value for column in range(1, 17)]
    columns = [1, 2, 3, 4, 5, None, None, None, None, 6, 7, 8, 9, 10, 11, 12] if source_format == "SCHOOL" else [1, 2, 3, 4, 5, 6, 7, 8, 9, None, None, None, None, None, None, None]
    return [_cell_value(sheet, row_no, column) if column else None for column in columns]


def _has_source_content(sheet, row_no: int, source_format: str) -> bool:
    column_count = 16 if source_format == "MERGED" else 12 if source_format == "SCHOOL" else 9
    return any(
        sheet.cell(row_no, column).value is not None and str(sheet.cell(row_no, column).value).strip()
        for column in range(1, column_count + 1)
    )


def _allocation_candidates(db: Session, raw: str | None, scope: str) -> list[dict[str, Any]]:
    if not raw or raw in {"无", "—"}:
        return []
    orgs = db.execute(select(SysOrgUnit).where(SysOrgUnit.status == "ENABLED")).scalars().all()
    aliases = db.execute(select(SysOrgAlias)).scalars().all()
    names: list[tuple[str, int | None, str, str]] = [(org.org_name, org.id, "STANDARD", "CONFIRMED") for org in orgs]
    names.extend((alias.alias_name, alias.org_id, "ALIAS", alias.confirm_status) for alias in aliases)
    candidates: list[dict[str, Any]] = []
    for name, org_id, match_type, confirm_status in sorted(names, key=lambda item: len(item[0]), reverse=True):
        start = raw.find(name)
        if start < 0:
            continue
        tail = raw[start + len(name):start + len(name) + 80]
        amount_match = re.search(r"[（(]([^）)]*)[）)]", tail)
        amount_raw = amount_match.group(1) if amount_match else None
        amount = parse_decimal_text(amount_raw) if amount_raw else None
        candidates.append(
            {
                "scope": scope,
                "raw_name": name,
                "match_type": match_type,
                "suggested_org_id": org_id,
                "confirm_status": confirm_status if org_id else "PENDING",
                "amount_raw": amount_raw,
                "amount": amount,
                "context": raw,
            }
        )
    return candidates


def _validate_row(db: Session, row: SourceImportRow) -> None:
    issues: list[dict[str, Any]] = []
    invalid = False
    if not row.direction_raw:
        issues.append({"field": "direction_raw", "message": "主要建设方向为空"})
        invalid = True
    if not row.indicator_raw:
        issues.append({"field": "indicator_raw", "message": "关键产出指标为空"})
        invalid = True
    type_obj = db.execute(select(DictIndicatorType).where(DictIndicatorType.type_name == row.indicator_type_raw)).scalars().first()
    unit = db.execute(select(DictIndicatorUnit).where(DictIndicatorUnit.unit_name == row.indicator_unit_raw)).scalars().first()
    if not type_obj:
        issues.append({"field": "indicator_type_raw", "message": "指标类型不在字典中"})
        invalid = True
    if not unit:
        issues.append({"field": "indicator_unit_raw", "message": "指标单位不在字典中"})
        invalid = True
    parsed_five = None
    parsed_2026 = None
    if unit:
        for field, raw in (("professional_five_year_raw", row.professional_five_year_raw), ("professional_2026_raw", row.professional_2026_raw)):
            parsed = parse_target(raw, unit)
            if field == "professional_five_year_raw":
                parsed_five = parsed
            else:
                parsed_2026 = parsed
            if parsed["parse_issue"]:
                issues.append({"field": field, "message": parsed["parse_issue"]})
        if unit.is_compound or unit.value_kind in {"RATIO", "NONE"}:
            issues.append({"field": "indicator_unit_raw", "message": "复合、比值或不适用单位须人工确认核算口径"})
    candidates = _allocation_candidates(db, row.professional_five_year_allocation_raw, "FIVE_YEAR")
    candidates += _allocation_candidates(db, row.professional_2026_allocation_raw, "YEAR_2026")
    if any(item["confirm_status"] != "CONFIRMED" or item["suggested_org_id"] is None for item in candidates):
        issues.append({"field": "allocation", "message": "学院别名或映射须人工确认"})
    for scope, parsed, label in (("FIVE_YEAR", parsed_five, "五年"), ("YEAR_2026", parsed_2026, "2026 年")):
        amounts = [item["amount"] for item in candidates if item["scope"] == scope and item["amount"] is not None]
        if parsed and parsed["value"] is not None and amounts:
            allocated = sum(amounts)
            if allocated != parsed["value"]:
                issues.append(
                    {
                        "field": "allocation",
                        "message": f"{label}括号分配合计 {allocated} 与目标 {parsed['value']} 不一致，禁止自动下发",
                        "allocated": allocated,
                        "target": parsed["value"],
                    }
                )
    from app.utils import json_value
    row.allocation_candidates_json = json_value(candidates)
    row.validation_json = json_value(issues)
    row.validation_status = "INVALID" if invalid else "NEED_CONFIRM" if issues else "IMPORTABLE"
    row.parse_status = "FAILED" if invalid else "NEED_CONFIRM" if issues else "PARSED"
    row.error_message = "；".join(str(issue["message"]) for issue in issues) or None


def import_excel(
    db: Session,
    *,
    plan: ConstructionPlan,
    user: SysUser,
    file_name: str,
    content: bytes,
    duplicate_mode: str = "SKIP",
    source_format: str = "MERGED",
) -> tuple[SourceImportFile, bool]:
    if plan.plan_status != "DRAFT":
        raise AppError(409, "IMPORT_PLAN_NOT_DRAFT", "仅草稿计划可导入源清单")
    if not content:
        raise AppError(422, "EMPTY_FILE", "上传文件为空")
    if source_format not in SOURCE_FORMATS:
        raise AppError(422, "INVALID_SOURCE_FORMAT", "不支持的源表格式")
    file_hash = hashlib.sha256(content).hexdigest()
    duplicate = db.execute(
        select(SourceImportFile)
        .where(SourceImportFile.plan_id == plan.id, SourceImportFile.file_hash == file_hash)
        .order_by(SourceImportFile.id.desc())
    ).scalars().first()
    if duplicate and duplicate_mode == "SKIP":
        return duplicate, True
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    except Exception as exc:
        raise AppError(422, "INVALID_EXCEL", "无法读取 Excel 工作簿") from exc
    settings = get_settings()
    object_key = f"imports/{uuid.uuid4().hex}.xlsx"
    path = settings.storage_dir / object_key
    path.write_bytes(content)
    batch = SourceImportFile(
        plan_id=plan.id,
        file_name=Path(file_name).name,
        object_key=object_key,
        file_hash=file_hash,
        sheet_count=len(workbook.worksheets),
        data_row_count=0,
        import_status="PREVIEW",
        imported_by=user.id,
        imported_at=utcnow(),
        created_by=user.id,
    )
    db.add(batch)
    db.flush()
    row_count = 0
    try:
        for sheet in workbook.worksheets:
            header = _find_header(sheet, source_format)
            if not header:
                continue
            header_row = header[0]
            current_direction = None
            for row_no in range(header_row + 1, sheet.max_row + 1):
                if not _has_source_content(sheet, row_no, source_format):
                    continue
                values = _source_values(sheet, row_no, source_format)
                data = {field: raw_text(value) for field, value in zip(ROW_FIELD_NAMES, values)}
                if source_format != "MERGED":
                    current_direction = data["direction_raw"] or current_direction
                    data["direction_raw"] = current_direction
                source_row = SourceImportRow(
                    import_file_id=batch.id,
                    sheet_name=sheet.title,
                    row_number=row_no,
                    created_by=user.id,
                    **data,
                )
                _validate_row(db, source_row)
                db.add(source_row)
                row_count += 1
        if row_count == 0:
            raise AppError(422, "NO_SOURCE_ROWS", "未找到符合所选源表格式的数据行")
        batch.data_row_count = row_count
        write_log(
            db,
            user,
            object_type="IMPORT_FILE",
            object_id=batch.id,
            object_name=batch.file_name,
            action_code="IMPORT_PREVIEW",
            plan_id=plan.id,
            after={"rows": row_count, "hash": file_hash, "source_format": source_format},
        )
        db.flush()
        return batch, False
    except Exception:
        if path.exists():
            path.unlink()
        raise


def confirm_import_row(
    db: Session,
    *,
    row: SourceImportRow,
    plan: ConstructionPlan,
    user: SysUser,
    regenerate: bool = False,
    overrides: dict[str, Any] | None = None,
) -> OutputIndicator:
    if row.indicator_id and not regenerate:
        return db.get(OutputIndicator, row.indicator_id)
    overrides = overrides or {}
    if row.validation_status == "INVALID" and not overrides:
        raise AppError(422, "IMPORT_ROW_INVALID", "不可导入行必须提供结构化修正后再确认")
    type_name = overrides.get("indicator_type_name", row.indicator_type_raw)
    unit_name = overrides.get("indicator_unit_name", row.indicator_unit_raw)
    type_obj = db.execute(select(DictIndicatorType).where(DictIndicatorType.type_name == type_name)).scalars().first()
    unit = db.execute(select(DictIndicatorUnit).where(DictIndicatorUnit.unit_name == unit_name)).scalars().first()
    if not type_obj or not unit:
        raise AppError(422, "IMPORT_DICTIONARY_MISMATCH", "修正后的指标类型或单位仍不在字典中")
    direction_name = overrides.get("direction_name", row.direction_raw)
    indicator_name = overrides.get("indicator_name", row.indicator_raw)
    if not direction_name or not indicator_name:
        raise AppError(422, "IMPORT_REQUIRED_FIELD", "方向和指标名称不能为空")
    direction = db.execute(
        select(ConstructionDirection).where(
            ConstructionDirection.plan_id == plan.id,
            ConstructionDirection.direction_name == direction_name,
        )
    ).scalars().first()
    if not direction:
        sequence = db.execute(
            select(ConstructionDirection).where(ConstructionDirection.plan_id == plan.id)
        ).scalars().all()
        source_seq = None
        try:
            source_seq = int(row.source_seq_raw) if row.source_seq_raw else None
        except ValueError:
            pass
        direction = ConstructionDirection(
            plan_id=plan.id,
            direction_code=f"DIR-{plan.id}-{len(sequence) + 1:03d}",
            direction_name=direction_name,
            source_seq=source_seq,
            sort_no=len(sequence) + 1,
            created_by=user.id,
        )
        db.add(direction)
        db.flush()
    existing_count = len(db.execute(select(OutputIndicator).where(OutputIndicator.plan_id == plan.id)).scalars().all())
    indicator_code = f"IND-{plan.id}-{existing_count + 1:04d}"
    five = parse_target(row.professional_five_year_raw, unit)
    school_five = parse_target(row.school_five_year_raw, unit)
    indicator = OutputIndicator(
        plan_id=plan.id,
        direction_id=direction.id,
        indicator_code=indicator_code,
        indicator_name=indicator_name,
        indicator_desc=row.indicator_raw or indicator_name,
        indicator_type_code=type_obj.type_code,
        unit_code=unit.unit_code,
        standard_unit_code=overrides.get("standard_unit_code"),
        unit_raw=row.indicator_unit_raw or unit.unit_name,
        target_direction=overrides.get("target_direction"),
        calc_method=overrides.get("calc_method"),
        statistical_rule=overrides.get("statistical_rule"),
        baseline_value=overrides.get("baseline_value"),
        baseline_year=overrides.get("baseline_year"),
        numerator_rule=overrides.get("numerator_rule"),
        denominator_rule=overrides.get("denominator_rule"),
        external_data_source=overrides.get("external_data_source"),
        source_import_row_id=row.id,
        parse_status="NEED_CONFIRM",
        created_by=user.id,
    )
    assign_target_fields(indicator, "professional_five_year_target", five)
    assign_target_fields(indicator, "school_five_year_target", school_five)
    db.add(indicator)
    db.flush()
    for year in range(plan.start_year, plan.end_year + 1):
        target = IndicatorAnnualTarget(
            indicator_id=indicator.id,
            target_year=year,
            confirm_status="PENDING",
            created_by=user.id,
        )
        if year == 2026:
            assign_target_fields(target, "professional_target", parse_target(row.professional_2026_raw, unit))
            assign_target_fields(target, "school_target", parse_target(row.school_2026_raw, unit))
        division = getattr(row, f"school_{year}_division_raw", None)
        target.school_division_raw = division
        db.add(target)
    row.indicator_id = indicator.id
    row.confirmed_by = user.id
    row.confirmed_at = utcnow()
    row.parse_status = "CONFIRMED"
    write_log(
        db,
        user,
        object_type="INDICATOR",
        object_id=indicator.id,
        object_name=indicator.indicator_name,
        action_code="IMPORT_CONFIRM_REGENERATE" if regenerate else "IMPORT_CONFIRM",
        plan_id=plan.id,
        related_order_type="IMPORT_ROW",
        related_order_id=row.id,
    )
    return indicator


def refresh_import_status(db: Session, batch: SourceImportFile) -> None:
    rows = db.execute(select(SourceImportRow).where(SourceImportRow.import_file_id == batch.id)).scalars().all()
    confirmed = sum(row.indicator_id is not None for row in rows)
    batch.import_status = "COMPLETED" if confirmed == len(rows) else "PARTIAL_CONFIRMED" if confirmed else "PREVIEW"


def initialization_stats(db: Session, plan: ConstructionPlan) -> dict[str, Any]:
    directions = db.execute(select(ConstructionDirection).where(ConstructionDirection.plan_id == plan.id)).scalars().all()
    indicators = db.execute(select(OutputIndicator).where(OutputIndicator.plan_id == plan.id)).scalars().all()
    indicator_ids = [item.id for item in indicators]
    annual = []
    if indicator_ids:
        annual = db.execute(select(IndicatorAnnualTarget).where(IndicatorAnnualTarget.indicator_id.in_(indicator_ids))).scalars().all()
    from app.models import CollegeTask

    tasks = db.execute(
        select(CollegeTask).join(OutputIndicator, OutputIndicator.id == CollegeTask.indicator_id).where(OutputIndicator.plan_id == plan.id)
    ).scalars().all()
    type_counts: dict[str, int] = {}
    unit_counts: dict[str, int] = {}
    for item in indicators:
        type_counts[item.indicator_type_code] = type_counts.get(item.indicator_type_code, 0) + 1
        unit_counts[item.unit_code] = unit_counts.get(item.unit_code, 0) + 1
    divisions = {
        str(year): sum(bool(item.school_division_raw) for item in annual if item.target_year == year)
        for year in range(plan.start_year, plan.end_year + 1)
    }
    return {
        "plan_id": plan.id,
        "direction_count": len(directions),
        "indicator_count": len(indicators),
        "annual_target_count": len(annual),
        "college_task_count": len(tasks),
        "indicator_types": type_counts,
        "unit_count": len(unit_counts),
        "units": unit_counts,
        "professional_five_year_empty": sum(item.professional_five_year_target_raw is None for item in indicators),
        "professional_2026_empty": sum(
            item.professional_target_raw is None for item in annual if item.target_year == 2026
        ),
        "school_division_nonempty": divisions,
    }
